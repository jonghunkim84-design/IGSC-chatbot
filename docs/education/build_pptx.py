"""사용자 교육 매뉴얼 PPTX 생성 (python-pptx).
사용: python docs/education/build_pptx.py      (필요: pip install python-pptx pillow)
텍스트가 상자를 넘치지 않도록 Malgun Gothic 글꼴 폭으로 줄바꿈을 미리 계산해 글자 크기를 자동으로 줄인다.
"""
import pathlib
import re

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = pathlib.Path(__file__).parent
OUT = HERE / "IGSC 챗봇 사용자 교육 매뉴얼.pptx"
FONT = "Malgun Gothic"
REG = ImageFont.truetype(r"C:\Windows\Fonts\malgun.ttf", 100)
BLD = ImageFont.truetype(r"C:\Windows\Fonts\malgunbd.ttf", 100)

TEAL = RGBColor(0x0F, 0x76, 0x6E)
TEAL_L = RGBColor(0xE2, 0xF3, 0xF1)
TEAL_XL = RGBColor(0xF0, 0xFD, 0xFA)
INK = RGBColor(0x1E, 0x29, 0x3B)
GRAY = RGBColor(0x64, 0x74, 0x8B)
LINE = RGBColor(0xCB, 0xD5, 0xE1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xB4, 0x53, 0x09)
AMBER_L = RGBColor(0xFE, 0xF3, 0xC7)
RED = RGBColor(0xB9, 0x1C, 0x1C)
RED_L = RGBColor(0xFE, 0xE2, 0xE2)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
warnings = []
page = [0]


# ── 글자 크기 계산 ──────────────────────────────────────────
def text_w(s, size, bold=False):
    return (BLD if bold else REG).getlength(s) * size / 100.0  # pt


def wrap_lines(s, size, width_pt, bold=False):
    """공백 기준 줄바꿈(너무 긴 단어는 글자 단위). 줄 수를 돌려준다."""
    s = s.replace("**", "")
    lines = 0
    for para in s.split("\n"):
        cur = ""
        n = 1
        for word in re.split(r"(?<=\s)", para):
            if text_w(cur + word, size, bold) <= width_pt:
                cur += word
            elif text_w(word, size, bold) > width_pt:
                for ch in word:
                    if text_w(cur + ch, size, bold) > width_pt:
                        n += 1
                        cur = ch
                    else:
                        cur += ch
            else:
                n += 1
                cur = word
        lines += n
    return lines


def block_height(paras, width_in, size):
    """paras: [(text, scale, bold, space_after_pt)] → 필요한 높이(pt)"""
    h = 0
    for i, (text, scale, bold, after) in enumerate(paras):
        sz = size * scale
        h += wrap_lines(text, sz, width_in * 72 - 14.4, bold) * sz * 1.2 + (after if i < len(paras) - 1 else 0)
    return h


def fit_size(paras, width_in, height_in, start, minimum=11):
    size = start
    while size > minimum and block_height(paras, width_in, size) > height_in * 72 - 6:
        size -= 0.5
    if block_height(paras, width_in, size) > height_in * 72 - 6:
        warnings.append(f"slide {page[0]}: 넘침 가능 ({paras[0][0][:20]!r}…)")
    return size


def set_run_font(run, size, bold=False, color=INK):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def add_rich(paragraph, text, size, bold=False, color=INK):
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        r = paragraph.add_run()
        if part.startswith("**") and part.endswith("**"):
            r.text = part[2:-2]
            set_run_font(r, size, True, TEAL if color == INK else color)
        else:
            r.text = part
            set_run_font(r, size, bold, color)


def textbox(slide, x, y, w, h, items, size=18, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullet=False, after=6, min_size=11):
    """items: 문자열 또는 (문자열, {bold, scale, color}) 목록"""
    norm = []
    for it in items:
        t, o = (it, {}) if isinstance(it, str) else it
        norm.append((t, o))
    paras = [(("• " if bullet else "") + t, o.get("scale", 1), o.get("bold", False), after) for t, o in norm]
    sz = fit_size(paras, w, h, size, min_size)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    for i, (t, o) in enumerate(norm):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        add_rich(p, ("• " if bullet else "") + t, sz * o.get("scale", 1), o.get("bold", False), o.get("color", color))
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


# ── 슬라이드 부품 ───────────────────────────────────────────
def new_slide(title, section=None):
    page[0] += 1
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 0.12, TEAL, shape=MSO_SHAPE.RECTANGLE)
    if section:
        textbox(s, 0.55, 0.28, 8, 0.35, [section], size=13, color=TEAL)
    textbox(s, 0.55, 0.58, 12.2, 0.8, [(title, {"bold": True})], size=30, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, 0.65, 1.38, 1.2, 0.05, TEAL, shape=MSO_SHAPE.RECTANGLE)
    textbox(s, 0.55, 7.05, 8, 0.3, ["IGSC 인증 문의 챗봇 · 사용자 교육"], size=10, color=GRAY)
    textbox(s, 11.8, 7.05, 1.0, 0.3, [str(page[0])], size=10, color=GRAY, align=PP_ALIGN.RIGHT)
    return s


def section_slide(num, title, sub):
    page[0] += 1
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 13.333, 7.5, TEAL, shape=MSO_SHAPE.RECTANGLE)
    textbox(s, 0.9, 2.2, 3, 1.2, [(str(num), {"bold": True})], size=72, color=RGBColor(0x99, 0xF6, 0xE4))
    textbox(s, 0.9, 3.5, 11.5, 1.0, [(title, {"bold": True})], size=40, color=WHITE)
    textbox(s, 0.9, 4.5, 11.5, 1.0, [sub], size=20, color=RGBColor(0xCC, 0xFB, 0xF1))
    return s


def bullets(slide, items, x=0.8, y=1.7, w=11.8, h=5.1, size=22, after=10):
    return textbox(slide, x, y, w, h, items, size=size, bullet=True, after=after)


def card(slide, x, y, w, h, title, body, fill=TEAL_XL, edge=TEAL, size=15):
    rect(slide, x, y, w, h, fill, edge)
    textbox(slide, x + 0.1, y + 0.08, w - 0.2, 0.55, [(title, {"bold": True})], size=size + 3, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
    body = [body] if isinstance(body, str) else body
    textbox(slide, x + 0.1, y + 0.68, w - 0.2, h - 0.78, body, size=size, after=4)


def callout(slide, x, y, w, h, text, kind="info", size=17):
    fill, edge, col = {"info": (TEAL_XL, TEAL, INK), "warn": (AMBER_L, AMBER, INK), "bad": (RED_L, RED, INK)}[kind]
    rect(slide, x, y, w, h, fill, edge)
    textbox(slide, x + 0.15, y + 0.05, w - 0.3, h - 0.1, [text] if isinstance(text, str) else text, size=size, color=col, anchor=MSO_ANCHOR.MIDDLE)


def flow(slide, steps, x=0.6, y=2.4, w=12.1, h=2.2, size=16, gap=0.28):
    n = len(steps)
    bw = (w - gap * (n - 1)) / n
    for i, (title, body) in enumerate(steps):
        bx = x + i * (bw + gap)
        rect(slide, bx, y, bw, h, TEAL_XL, TEAL)
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(bx + 0.12), Inches(y + 0.12), Inches(0.42), Inches(0.42))
        circ.fill.solid()
        circ.fill.fore_color.rgb = TEAL
        circ.line.fill.background()
        tf = circ.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = str(i + 1)
        set_run_font(r, 14, True, WHITE)
        textbox(slide, bx + 0.55, y + 0.08, bw - 0.6, 0.55, [(title, {"bold": True})], size=size + 1, color=TEAL, anchor=MSO_ANCHOR.MIDDLE)
        textbox(slide, bx + 0.05, y + 0.7, bw - 0.1, h - 0.8, [body], size=size - 1, after=2)
        if i < n - 1:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(bx + bw + 0.03), Inches(y + h / 2 - 0.12), Inches(gap - 0.06), Inches(0.24))
            ar.fill.solid()
            ar.fill.fore_color.rgb = TEAL
            ar.line.fill.background()


def table(slide, x, y, w, cols, rows, size=14, header=True, max_h=5.2, min_size=10, first_bold=True):
    """cols: 열 너비 비율 목록. rows[0] 이 머리글."""
    total = sum(cols)
    widths = [w * c / total for c in cols]

    def heights(sz):
        hs = []
        for ri, row in enumerate(rows):
            lines = max(wrap_lines(c, sz, widths[ci] * 72 - 14.4, bold=(ri == 0 or (ci == 0 and first_bold))) for ci, c in enumerate(row))
            hs.append(lines * sz * 1.25 + 12)
        return hs

    sz = size
    while sz > min_size and sum(heights(sz)) / 72 > max_h:
        sz -= 0.5
    hs = heights(sz)
    if sum(hs) / 72 > max_h + 0.05:
        warnings.append(f"slide {page[0]}: 표 높이 초과 ({sum(hs)/72:.1f}in > {max_h}in)")
    shape = slide.shapes.add_table(len(rows), len(cols), Inches(x), Inches(y), Inches(w), Emu(int(sum(hs) / 72 * 914400)))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is not None:
        sid.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"  # 스타일 없음(격자)
    for ci, cw in enumerate(widths):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Emu(int(hs[ri] / 72 * 914400))
        for ci, text in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL if (header and ri == 0) else (WHITE if ri % 2 else RGBColor(0xF8, 0xFA, 0xFC))
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            if header and ri == 0:
                add_rich(p, text.replace("**", ""), sz, True, WHITE)
            else:
                add_rich(p, text, sz, bold=(ci == 0 and first_bold))
            for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                tcPr = cell._tc.get_or_add_tcPr()
                ln = tcPr.find(qn(side))
                if ln is not None:
                    tcPr.remove(ln)
                ln = tcPr.makeelement(qn(side), {"w": "9525"})
                sf = ln.makeelement(qn("a:solidFill"), {})
                clr = sf.makeelement(qn("a:srgbClr"), {"val": "CBD5E1"})
                sf.append(clr)
                ln.append(sf)
                tcPr.insert(0, ln) if side == "a:lnL" else tcPr.insert(1, ln)
    return shape


# ═══════════════════════ 슬라이드 ═══════════════════════
# 1 표지
page[0] += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, TEAL, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0.9, 2.55, 1.6, 0.08, RGBColor(0x99, 0xF6, 0xE4), shape=MSO_SHAPE.RECTANGLE)
textbox(s, 0.8, 2.75, 11.8, 1.3, [("IGSC 인증 문의 챗봇", {"bold": True})], size=48, color=WHITE)
textbox(s, 0.8, 3.95, 11.8, 0.9, ["사용자 교육 매뉴얼"], size=34, color=RGBColor(0xCC, 0xFB, 0xF1))
textbox(s, 0.8, 5.4, 11.8, 0.9, ["운영 담당자(관리자)용 · 기준일 2026-10-04", "챗봇 igsc-chatbot.vercel.app/chat · 관리자 igsc-chatbot.vercel.app/admin"], size=16, color=RGBColor(0xCC, 0xFB, 0xF1))

# 2 교육 안내
s = new_slide("교육 안내")
card(s, 0.6, 1.7, 3.9, 2.4, "대상", ["관리자 계정을 받는 운영 담당자", "(관리자 최대 3명)"], size=17)
card(s, 4.7, 1.7, 3.9, 2.4, "준비물", ["관리자로 등록된 이메일", "컴퓨터와 인터넷 브라우저", "(로그인 메일의 링크는 로그인을 요청한 같은 브라우저에서 열기)"], size=15)
card(s, 8.8, 1.7, 3.9, 2.4, "소요", ["약 2시간 30분", "(실습 포함)"], size=17)
textbox(s, 0.6, 4.3, 12, 0.5, [("교육 목표", {"bold": True})], size=20, color=TEAL)
bullets(s, ["챗봇이 어떤 규칙으로 답하는지 설명할 수 있다", "FAQ를 만들고, 확인하고, 승인할 수 있다", "챗봇이 답하지 못한 질문을 처리하고, 파일에서 FAQ 초안을 만들어 확인할 수 있다", "매일·매주·매달 해야 할 일을 안다"], y=4.8, h=2.1, size=18, after=4)

# 3 진행 순서
s = new_slide("진행 순서 (예시)")
table(
    s, 0.8, 1.75, 11.7, [1.3, 6, 1.4],
    [["시간", "내용", "장"], ["15분", "챗봇이 무엇이고 어떤 규칙으로 답하는지", "1장"], ["10분", "고객이 보는 챗봇 화면 시연", "2장"], ["15분", "관리자 화면 로그인과 메뉴", "3장"],
     ["60분", "핵심 업무 따라 하기 (FAQ 만들기 · 미답변 처리 · 문서로 초안 만들기 · 재검토) + 실습", "4장"], ["15분", "FAQ를 잘 쓰는 방법과 승인 전 체크리스트", "5장"], ["15분", "운영 루틴과 운영 상태 화면", "6장"],
     ["20분", "정리 · 질의응답 · 문제가 생겼을 때", "7·8장"]],
    size=18, max_h=4.2,
)
callout(s, 0.8, 6.0, 11.7, 0.85, "**교육 중 주의** — 지금은 승인된 FAQ가 한 건도 없습니다(0건). 실습에서 **승인한 FAQ는 실제 챗봇에 바로 나갑니다.**\n실습 FAQ는 질문 앞에 [실습]을 붙이고, 끝나면 꼭 초안으로 되돌리거나 삭제하세요.", "warn", size=15)

s = new_slide("처음 보시는 분을 위한 용어 풀이")
table(s, 0.5, 1.65, 12.3, [1.6, 8], [
    ["용어", "쉬운 설명"],
    ["FAQ", "\"자주 묻는 질문\"과 미리 정해 둔 \"답변\"을 한 쌍으로 묶은 것. **챗봇은 FAQ에 있는 내용으로만 답합니다.**"],
    ["초안 / 승인", "**초안**: 작성·확인 중이라 고객에게 보이지 않는 상태. **승인**: 사람이 확인을 마쳐서 챗봇이 쓰기 시작하는 상태"],
    ["다른 표현", "같은 질문을 사람마다 다르게 말하는 것을 모아 둔 것. 예) \"비건 인증 비용은?\" = \"비건인증 얼마\""],
    ["담당자 이관", "챗봇이 답하지 않고 \"담당자가 확인해 드립니다\"라고 안내하며 사람에게 넘기는 것"],
    ["미답변 질문", "챗봇이 답하지 못해 담당자에게 넘긴 질문 목록. **새 FAQ를 만드는 출발점**"],
    ["보완 필요", "답변을 아직 정하지 못해 **비워 둔 FAQ**. 승인할 수 없고 챗봇도 쓰지 않음"],
    ["문서 보관함", "인증 안내서 같은 **파일을 올려 두는 곳**. 파일에서 FAQ 초안을 자동으로 만들 수 있음"],
], size=15, max_h=5.2)

# ── 1장
section_slide(1, "챗봇이 일하는 방식", "챗봇의 규칙과 질문이 처리되는 과정")

s = new_slide("무엇인가", "1. 챗봇이 일하는 방식")
textbox(s, 0.6, 1.65, 12.1, 1.1, ["고객이 홈페이지 챗봇에 인증에 대해 물어보면(가능 여부, 비용, 기간, 절차, 서류 등), 사람이 미리 확인해 승인한 **FAQ의 내용만** 가지고 답합니다. FAQ에 없거나 답하면 안 되는 질문은 **억지로 답하지 않고 담당자에게 연결**합니다. (챗봇이 지어내서 답하는 일을 막기 위한 구조입니다.)"], size=20)
card(s, 0.6, 3.0, 3.9, 3.3, "챗봇 화면", ["홈페이지 오른쪽 아래 둥근 버튼", "고객의 질문에 답합니다", "사용하는 사람: 고객"], size=17)
card(s, 4.7, 3.0, 3.9, 3.3, "관리자 화면 (/admin)", ["FAQ 만들기·승인", "미답변 질문 처리", "문서 올리기·초안 만들기", "대화 기록·통계 보기", "사용하는 사람: 운영 담당자"], size=16)
card(s, 8.8, 3.0, 3.9, 3.3, "데이터 저장소", ["FAQ, 대화 기록", "문서 파일 등을 안전하게 보관", "(시스템이 알아서 관리)"], size=17)

s = new_slide("챗봇이 반드시 지키는 5가지 규칙", "1. 챗봇이 일하는 방식")
cw = 2.3
items = [
    ("① 승인된 FAQ로만", "맞는 FAQ가 없으면 답을 지어내지 않고 담당자에게 연결"),
    ("② 컨설팅은 하지 않음", "\"어떻게 하면 통과하나요?\"에는 방법을 알려 주지 않음. 인증기관이 코치하면 심사가 공정하지 않기 때문(ISO 17065)"),
    ("③ 비용·기간은 확정값만", "추측하지 않고, 항상 \"참고용이며 담당자 검토 후 확정\" 문구를 붙임"),
    ("④ 인증 가능 여부는 확답하지 않음", "비슷한 제품의 인증 사례만 참고로 보여 주고, 최종 판단은 담당자"),
    ("⑤ 모든 대화를 기록", "답하지 못한 질문은 관리자에게 모임"),
]
for i, (t, b) in enumerate(items):
    card(s, 0.6 + i * (cw + 0.15), 1.8, cw, 3.4, t, b, size=14)
callout(s, 0.6, 5.5, 12.1, 1.1, "이 규칙들은 사람이 신경 써서 지키는 것이 아니라 **시스템이 스스로 지킵니다.** 운영자는 FAQ를 잘 만드는 데 집중하면 됩니다.", "info", size=18)

s = new_slide("운영의 핵심: 질문이 FAQ가 되는 순환", "1. 챗봇이 일하는 방식")
flow(s, [("고객이 질문", "챗봇이 승인된 FAQ로 답변"), ("답하지 못함", "미답변 질문 목록에 쌓임"), ("관리자가 처리", "FAQ로 만들고 승인"), ("바로 반영", "다음 고객부터는 바로 답변")], y=2.0, h=2.4, size=17)
callout(s, 0.6, 5.0, 12.1, 1.3, ["**챗봇이 똑똑해지는 방법 = 승인된 FAQ를 늘리고 다듬는 것**", "운영 담당자의 일은 이 순환을 꾸준히 돌리는 것입니다."], "info", size=22)

s = new_slide("질문 하나가 처리되는 과정", "1. 챗봇이 일하는 방식")
flow(s, [("분류", "질문 종류 구분: 인증 문의 / 컨설팅 / 불만 / 관계없는 질문"), ("갈림길", "컨설팅·불만·관계없는 질문은 정해 둔 안내 문구로 끝"), ("FAQ 고르기", "승인된 FAQ 중 관련된 것을 최대 5개 + 모든 인증에 해당하는 공통 FAQ"), ("답변 쓰기", "고른 FAQ 내용만으로 답변. 출처 링크는 자동으로 붙음"), ("기록", "대화 기록에 남기고, 답 못한 질문은 미답변 목록에")], y=1.9, h=3.0, size=15, gap=0.25)
callout(s, 0.6, 5.3, 12.1, 1.0, "컨설팅·불만·관계없는 질문은 AI가 새로 답을 만들지 않고, **미리 정해 둔 안내 문구**만 나갑니다.", "info", size=18)

s = new_slide("처리 결과 7가지와 운영자가 할 일", "1. 챗봇이 일하는 방식")
table(s, 0.5, 1.65, 12.3, [1.5, 3.2, 3.4, 3.6], [
    ["처리 결과", "어떤 경우인가요?", "고객에게 나가는 것", "운영자가 할 일"],
    ["답변 완료", "승인 FAQ로 답함", "답변 + 출처 링크", "가끔 정확성 확인"],
    ["부분 안내", "질문 중 일부만 FAQ로 답할 수 있음", "답할 수 있는 부분까지 답하고 나머지는 \"담당자가 확인\"", "**미답변 질문**에서 빠진 부분을 FAQ로 추가"],
    ["담당자 이관", "맞는 FAQ 없음 / 보완 필요 FAQ / 시스템 오류 / 외국어", "\"담당자 확인이 필요합니다\" + 연락처", "**미답변 질문**에서 FAQ로 만들기 (가장 중요)"],
    ["인증 종류 되묻기", "어떤 인증인지 알 수 없음 (예: \"얼마나 걸려요?\")", "\"어떤 인증인가요?\" + 선택 버튼", "없음 (고객이 버튼을 누르면 이어짐)"],
    ["컨설팅 차단", "통과하는 방법, 기준 맞추는 법을 물음", "공평성 안내 + 외부 컨설팅 기관 안내", "없음. 이런 질문에 답하는 FAQ는 **만들지 않기**"],
    ["불만 접수", "불만·이의제기·항의", "접수 방법 + 연락처", "**담당자가 직접 대응**"],
    ["범위 밖", "인증과 관계없는 질문 (날씨, 잡담 등)", "\"인증 문의 전용 챗봇\" 안내", "없음"],
], size=14, max_h=5.2)

s = new_slide("시스템이 알아서 붙여 주는 문구", "1. 챗봇이 일하는 방식")
table(s, 0.8, 1.8, 11.7, [4.6, 4], [
    ["자동으로 붙는 것", "언제 붙나요?"],
    ["출처 링크", "FAQ에 출처 URL이 있을 때"],
    ["\"비용·기간은 참고용이며 담당자 검토 후 확정됩니다\"", "비용·기간 FAQ로 답할 때"],
    ["\"인증 가능 여부는 담당자가 검토한 뒤 확정됩니다\"", "제품의 인증 가능 여부를 물을 때"],
    ["비슷한 제품의 인증 사례", "가능 여부를 물었고, 공개에 동의한 사례가 있을 때"],
], size=18, max_h=3.4)
callout(s, 0.8, 5.4, 11.7, 1.0, "이런 문구는 FAQ 답변에 **직접 쓰지 않아도** 됩니다.", "info", size=20)

s = new_slide("'초안'과 '승인' — 가장 중요한 개념", "1. 챗봇이 일하는 방식")
bullets(s, [
    "**승인된 FAQ만** 챗봇이 사용합니다. **초안은 고객에게 절대 보이지 않는** 작업 공간입니다.",
    "승인은 **사람이 직접 버튼을 눌러야만** 됩니다. 파일에서 자동으로 만든 초안도 예외가 없습니다.",
    "저장하면 챗봇 답변에 **바로** 반영됩니다. (기다리거나 따로 배포할 필요 없음)",
    "누가 언제 승인했는지 이메일과 시각이 기록됩니다. (FAQ 목록의 '승인자' 칸)",
    "**'보완 필요'(담당자 입력 필요)**: 답변이 아직 정해지지 않아 비워 둔 FAQ. 답변을 채우고 표시를 해제해야 승인할 수 있고, 그 전에는 챗봇이 쓰지 않습니다.",
], size=20, after=12)

# ── 2장
section_slide(2, "고객이 보는 챗봇 화면", "챗봇은 고객에게 이렇게 보입니다")

s = new_slide("챗봇 사용법과 좋은 질문", "2. 고객이 보는 챗봇 화면")
card(s, 0.6, 1.7, 5.9, 2.5, "시작하기", ["홈페이지 오른쪽 아래 동그란 버튼 클릭", "추천 질문 버튼을 누르거나, 질문을 쓰고 '전송'", "PC는 작은 대화창, 스마트폰은 전체 화면"], size=16)
card(s, 6.8, 1.7, 5.9, 2.5, "입력 규칙", ["인증 이름을 함께 쓰면 가장 정확", "한 번에 하나의 질문, 500자 이내, 한국어", "연락처 같은 개인정보는 입력하지 않도록 안내"], size=16)
table(s, 0.6, 4.45, 12.1, [1, 1], [["좋은 질문", "아쉬운 질문"], ["비건 인증 비용이 얼마인가요?", "얼마예요?"], ["식품 비건 인증은 어떤 원칙으로 심사하나요?", "심사는 어떻게 해요?"]], size=17, max_h=1.9, first_bold=False)
textbox(s, 0.6, 6.45, 12.1, 0.5, ["인증 이름이 없으면 \"어떤 인증에 대한 문의이신가요?\"라고 되묻고, 버튼을 누르면 이어서 답합니다."], size=15, color=GRAY)

s = new_slide("시연 순서 (강사용)", "2. 고객이 보는 챗봇 화면")
flow(s, [("비용을 묻는다", "안내 문구가 붙는 것 확인"), ("컨설팅 질문", "\"어떻게 하면 통과하나요?\" → 정중한 거절"), ("모호한 질문", "\"얼마나 걸려요?\" → 어떤 인증인지 되묻기"), ("관계없는 질문", "→ \"인증 문의 전용\" 안내"), ("불만을 말한다", "→ 담당자 연결 안내")], y=1.9, h=2.6, size=16, gap=0.25)
callout(s, 0.6, 5.0, 12.1, 1.5, ["같은 사람이 짧은 시간에 질문을 너무 많이 보내면 \"잠시 후 다시 시도해 주세요\"가 나옵니다.", "(사용량과 AI 비용이 갑자기 늘어나는 것을 막는 장치: 한 사람당 1분에 15번, 1시간에 120번까지)"], "info", size=17)

# ── 3장
section_slide(3, "관리자 화면 시작하기", "로그인과 메뉴 둘러보기")

s = new_slide("로그인 (비밀번호 없이 이메일로)", "3. 관리자 화면")
flow(s, [("접속", "브라우저에서 /admin 주소로 들어가 등록된 이메일 입력"), ("링크 받기", "'로그인 링크 받기' 버튼 클릭"), ("메일 확인", "받은 메일의 링크를 요청한 같은 브라우저에서 클릭")], y=1.9, h=2.4, size=17, gap=0.4)
bullets(s, ["등록되지 않은 이메일은 메일이 오지 않습니다 (화면에는 똑같은 안내가 나옵니다)", "메일이 안 오면 스팸함을 확인하세요", "보안을 위해 로그인 요청을 짧은 시간에 반복하면 잠시 막힙니다 (1시간에 5번)"], y=4.6, h=2.2, size=18, after=8)

s = new_slide("화면 위쪽 메뉴", "3. 관리자 화면")
table(s, 0.6, 1.65, 12.1, [1.8, 6.5], [
    ["메뉴", "여기서 하는 일"],
    ["대시보드", "오늘·이번 주 문의 수, 답변 성공률, 많이 묻는 질문 Top 10, 처리할 미답변 질문 수를 한눈에 봄"],
    ["FAQ", "FAQ 찾기, 새로 만들기, 고치기, 여러 건 한꺼번에 승인·삭제"],
    ["미답변 질문", "챗봇이 답하지 못한 질문을 FAQ로 만들거나 '처리 완료'로 표시"],
    ["대화 로그", "고객과 챗봇의 대화를 날짜·처리 결과별로 다시 보기"],
    ["문서 보관함", "파일 올리기, 파일에서 FAQ 초안 만들기, 초안 확인"],
    ["인증 종류", "인증 이름(비건, 유기농 등) 추가·수정"],
    ["운영 상태", "챗봇이 지금 정상적으로 일하는지 확인"],
], size=16, max_h=4.8)

# ── 4장
section_slide(4, "핵심 업무 따라 하기", "FAQ 만들기 · 미답변 처리 · 문서로 초안 만들기 · 재검토")

s = new_slide("FAQ 새로 만들기", "4. 핵심 업무")
card(s, 0.6, 1.7, 6.1, 4.9, "입력 항목", [
    "**질문**: 고객이 물어볼 법한 문장 (예: 비건 인증 비용은 얼마인가요?)",
    "**다른 표현**: 같은 질문을 다르게 말하는 문장, 한 줄에 하나씩 (예: 비건인증 얼마)",
    "**답변**: 확정된 내용만 쓰기",
    "**인증 종류, 카테고리**: 절차 / 비용 / 기간 / 서류 / 대상·범위 / 갱신 중 선택",
    "**출처 URL**: 근거 홈페이지 주소. 고객에게 '출처' 링크로 보임",
], size=15)
card(s, 6.9, 1.7, 5.8, 3.1, "저장", ["위쪽 메뉴 FAQ → **+ FAQ 등록**", "상태를 **승인**으로 저장하면 챗봇이 바로 사용", "더 다듬을 게 있으면 **초안**으로 저장해 두기"], size=16)
callout(s, 6.9, 5.0, 5.8, 1.6, "질문을 쓰는 중에 **\"비슷한 FAQ가 있습니다\"** 경고가 뜰 수 있습니다. 같은 내용이면 새로 만들지 말고 기존 FAQ를 고치세요. (경고일 뿐 저장은 막지 않습니다)", "warn", size=15)

s = new_slide("FAQ 고치기 · 챗봇에서 내리기 · 삭제", "4. 핵심 업무")
card(s, 0.6, 1.7, 3.9, 4.3, "수정", ["목록 위쪽 검색칸으로 찾기", "질문 클릭 → 고치기 → **저장**", "저장하면 챗봇에 바로 반영"], size=16)
card(s, 4.7, 1.7, 3.9, 4.3, "챗봇에서 내리기", ["상태를 **초안**으로 바꿔 저장", "지우는 것이 아니라 잠시 내리는 것", "다시 승인하면 바로 돌아옴"], size=16)
card(s, 8.8, 1.7, 3.9, 4.3, "삭제", ["**초안 상태만** 삭제 가능", "수정 화면 아래 **삭제** 버튼 (확인 창이 한 번 뜸)", "**승인된 FAQ는 삭제되지 않음**", "삭제하면 되돌릴 수 없음"], size=16, fill=RED_L, edge=RED)
callout(s, 0.6, 6.15, 12.1, 0.7, "챗봇에서 내리고 싶을 때는 삭제보다 **초안으로 바꾸기**가 안전합니다.", "info", size=17)

s = new_slide("여러 건을 한꺼번에 승인·삭제하기", "4. 핵심 업무")
flow(s, [("고르기", "각 FAQ 왼쪽 체크 칸 클릭 (맨 위 체크 칸 = 현재 페이지 전체)"), ("버튼 누르기", "목록 아래 **선택 항목 승인** 또는 **선택 항목 삭제**"), ("결과 확인", "처리한 건수와 건너뛴 건수가 화면에 표시")], y=1.8, h=2.3, size=16, gap=0.4)
bullets(s, [
    "**승인**: 답변이 비어 있거나 '보완 필요'가 켜진 항목은 **자동으로 건너뜁니다**",
    "**삭제**: 초안만 삭제됩니다. 승인된 FAQ가 섞여 있으면 건너뜁니다 (확인 창에 선택한 건수가 표시됨)",
    "승인하면 FAQ 목록의 **승인자** 칸에 내 이메일이 남습니다",
], y=4.4, h=2.4, size=18, after=8)

s = new_slide("미답변 질문 처리 — 매일 하는 일 (가장 중요)", "4. 핵심 업무")
flow(s, [("미답변 질문 열기", "위쪽 메뉴 또는 대시보드의 '미처리 미답변 질문' 숫자 클릭"), ("직접 등록", "그 질문이 미리 채워진 FAQ 등록 화면이 열림"), ("채우고 저장", "답변·인증 종류·카테고리를 쓰고 저장"), ("자동 처리 완료", "그 미답변 질문이 '처리 완료'로 바뀜")], y=1.8, h=2.4, size=16, gap=0.3)
bullets(s, [
    "**부분 안내** 표시: 챗봇이 일부만 답하고 나머지를 담당자에게 넘긴 질문 → 빠진 부분을 FAQ로 채워 주세요",
    "인증을 되물어 고객이 고른 경우는 \"인증 절차가 어떻게 되나요? / 환경성적표지(EPD)\"처럼 처음 질문과 고른 인증이 함께 기록됩니다",
    "FAQ가 필요 없는 질문(잡담 등)은 **처리 완료** 버튼만 누르세요. 실수했다면 '처리 완료' 탭에서 미처리로 되돌릴 수 있습니다",
], y=4.5, h=2.4, size=17, after=8)

s = new_slide("파일에서 FAQ 초안 만들기: 전체 흐름", "4. 핵심 업무")
flow(s, [("파일 올리기", "문서 보관함에서 업로드"), ("글자 읽어오기", "자동. '추출 완료'가 되면 OK"), ("초안 만들기", "파일에서 근거를 찾아 초안 작성"), ("사람이 확인", "원문과 비교"), ("승인", "사람이 버튼 클릭 → 챗봇에 반영")], y=1.8, h=2.0, size=15, gap=0.25)
bullets(s, [
    "**파일은 답변의 근거가 아닙니다.** 챗봇은 파일을 직접 읽지 않고 승인된 FAQ로만 답합니다. 파일은 초안을 쓰는 재료입니다",
    "만들어진 초안은 **항상 '초안'** 상태이고, 승인 전에는 고객에게 보이지 않습니다",
    "파일에 근거가 없으면 답을 지어내지 않고, 답변이 빈 '보완 필요' 초안으로 남습니다",
    "숫자(비용·기간)와 근거 문장은 원문과 **글자 그대로 비교**해서, 다르면 자동으로 버립니다",
    "컨설팅에 해당하는 내용은 초안으로 만들지 않습니다",
], y=4.1, h=2.8, size=16, after=5)

s = new_slide("파일 올리기 규칙", "4. 핵심 업무")
table(s, 0.5, 1.65, 12.3, [1.7, 6.5], [
    ["항목", "규칙"],
    ["올리기", "한 번에 최대 20개, 파일 하나당 20MB. **인증 종류**와 **분류**를 함께 선택"],
    ["분류", "절차 / 비용·수수료 / 양식 / 규정·정책 중 선택. **'기타(초안 제외)'를 고르면 초안 재료로 쓰이지 않음**. 내부 설명서 같은 자료는 올리지 않기"],
    ["파일 종류", "PDF, 워드(DOCX), 엑셀(XLSX), 파워포인트(PPTX), CSV, TXT, MD. **한글(HWP)은 불가** → PDF로 저장해서 올리기"],
    ["스캔한 PDF", "사진이라 글자를 읽을 수 없어 '추출 실패'. 글자가 있는 원본(워드, 파워포인트 등)으로 다시 올리기"],
    ["같은 파일", "내용이 똑같은 파일은 올릴 수 없음. 이름만 같고 내용이 다르면 **개정판으로 올리기**(기본) / 별도 문서 / 취소 중 선택"],
    ["자료 형식", "일반 설명 글은 **설명 자료**(기본), 질문·답이 정리된 Q&A 파일은 **질문·답변 자료**. Q&A 자료는 원문을 그대로 가져오며 '형식 확인'에서 미리보기 3건을 보고 진행"],
], size=15, max_h=5.2)

s = new_slide("초안 만들기와 확인하기", "4. 핵심 업무")
card(s, 0.6, 1.7, 5.9, 5.0, "초안 만들기", [
    "해당 파일 줄의 **FAQ 초안 만들기** 클릭 (여러 파일은 체크 후 맨 위 버튼)",
    "시간이 걸리며 **조각 n/m**으로 진행 상황 표시. 화면을 닫으면 멈추지만 **이어서 진행**으로 계속 가능",
    "파일 하나에서 최대 **50건**",
    "끝나면 **초안 N건 생성 · 탈락 M건** — '탈락 목록 보기'에서 이유 확인",
    "자동 탈락: 근거가 원문에 없음, 숫자가 다름, 컨설팅, 다른 기관 안내, 이미 있는 FAQ와 같음",
], size=14)
card(s, 6.7, 1.7, 6.0, 5.0, "확인하기 (파일 줄의 '검수하기' 버튼)", [
    "카드에서 질문·답변을 바로 고칠 수 있음. **원문 보기**를 펼치면 근거 문장이 **노란색**으로 표시되니 원문과 비교",
    "인증 종류·카테고리도 카드마다 바꿀 수 있음",
    "**승인**(고친 내용 저장 + 승인, 바로 반영) / **보류**(초안으로 남김) / **삭제**(초안만)",
    "화면 아래 **일괄·전체 승인** (확인 창이 한 번 뜸). 키보드: ↑↓ 이동, A 승인, S 보류, D 삭제",
    "**꼭 확인**: 숫자(비용·기간), 대상과 조건, 원문에 없는 문장이 섞이지 않았는지",
], size=14)

s = new_slide("파일 내용이 바뀌었을 때 (재검토)", "4. 핵심 업무")
flow(s, [("개정판 올리기", "문서 보관함에서 해당 파일의 '개정판 올리기'"), ("재검토 표시", "이전 파일로 만든 승인 FAQ에 '재검토 필요'가 자동으로 붙음"), ("비교", "'새 문서로 초안 다시 만들기' → 기존 답변과 새 초안이 나란히 보임"), ("처리", "새 초안이 맞으면 채워서 저장 / 바꿀 게 없으면 '변경 없음'")], y=1.8, h=2.6, size=15, gap=0.3)
bullets(s, [
    "재검토하는 동안에도 **기존 답변은 계속 고객에게 나갑니다** (서비스가 끊기지 않음)",
    "대시보드의 '재검토 필요 N건' 카드, FAQ 목록의 '재검토 필요만' 체크로 모아서 볼 수 있습니다",
    "파일을 삭제한 경우에도 같은 표시가 붙습니다",
], y=4.7, h=2.1, size=18, after=8)

s = new_slide("인증 종류 추가 · 대시보드와 대화 로그 읽기", "4. 핵심 업무")
card(s, 0.6, 1.7, 5.9, 5.0, "새 인증 종류 추가 (인증 종류 메뉴)", [
    "코드(영문 소문자·숫자·하이픈. **한 번 만들면 바꿀 수 없음**), 한글 이름, 영문 이름, 분류, 정렬 순서를 쓰고 **추가**",
    "파일 올리기·FAQ 등록 화면의 선택 목록에 바로 나타남",
    "숨겨도(비활성화) 기존 FAQ·문서는 그대로, 새로 등록할 때 선택 목록에서만 빠짐",
    "**공통**은 모든 인증에 해당하는 특별 항목이라 **고정**이며 바꿀 수 없음",
], size=14)
card(s, 6.7, 1.7, 6.0, 5.0, "대시보드·대화 로그", [
    "**답변 성공률** = 답변 완료 ÷ (답변 완료 + 담당자 이관). '챗봇이 답했어야 하는 질문 중 실제로 답한 비율'",
    "**많이 묻는 질문 Top 10**: 최근 30일간 가장 많이 들어온 질문. '담당자 이관'이 많은 질문이 **FAQ로 만들 1순위**",
    "**대화 로그**: 대화를 펼치면 챗봇이 한 말과 근거 FAQ가 보임. 답이 이상하면 근거 FAQ를 열어 고치기",
    "**불만 접수** 대화는 담당자가 직접 읽고 대응",
], size=14)

# ── 5장
section_slide(5, "FAQ를 잘 쓰는 방법", "챗봇의 답변 품질은 FAQ에 달려 있습니다")

s = new_slide("FAQ 쓰기 5가지 요령", "5. FAQ를 잘 쓰는 방법")
items = [
    ("① 하나에 하나", "FAQ 하나에는 질문 하나만. 여러 질문을 섞으면 챗봇이 헷갈림"),
    ("② 인증 이름 넣기", "이름이 같은 인증은 식품 비건 / 화장품 비건처럼 분야까지"),
    ("③ 확정된 사실만", "비용·기간은 확정값 + 기준일 + 달라지는 조건. 모르면 비워 두기"),
    ("④ 다른 표현 충분히", "대화 로그에서 고객이 실제로 쓴 말을 그대로 추가"),
    ("⑤ 출처 주소 넣기", "고객이 원문을 직접 확인할 수 있어 신뢰도가 올라감"),
]
for i, (t, b) in enumerate(items):
    card(s, 0.6 + i * 2.45, 1.8, 2.3, 4.2, t, b, size=15)

s = new_slide("좋은 FAQ와 고쳐야 할 FAQ", "5. FAQ를 잘 쓰는 방법")
table(s, 0.6, 1.65, 12.1, [1, 3, 3], [
    ["", "좋은 FAQ", "고치면 좋은 FAQ"],
    ["질문", "식품 비건(Vegan) 인증의 원칙은 무엇인가요?", "비건 관련 (너무 넓어서 무엇을 묻는지 알 수 없음)"],
    ["답변", "동물 및 동물 유래 성분 금지, 동물 실험 금지, 교차 오염 금지의 세 가지 원칙입니다.", "보통 3~4주 걸립니다 (근거 없는 짐작)"],
], size=16, max_h=2.4)
callout(s, 0.6, 4.3, 12.1, 2.4, ["**FAQ에 넣으면 안 되는 것**", "• 컨설팅 내용 (통과하는 방법, 기준 맞추는 법, 준비·개선 조언)", "• 짐작이나 일반 상식 (\"보통\", \"대략\", \"일반적으로\" 같은 표현)", "• 개인정보, 다른 기관의 정보"], "bad", size=18)

s = new_slide("승인 버튼 누르기 전 체크리스트", "5. FAQ를 잘 쓰는 방법")
bullets(s, [
    "☐  질문이 하나의 주제이고, 어떤 인증인지 드러난다",
    "☐  답변의 모든 내용이 공식 자료 또는 담당자가 확정한 값과 같다",
    "☐  비용·기간이면 기준일과 달라지는 조건을 적었다",
    "☐  컨설팅처럼 읽히는 표현이 없다",
    "☐  출처 주소와 '다른 표현'이 있다",
], y=1.8, h=3.6, size=24, after=14)
callout(s, 0.8, 5.5, 11.7, 1.1, "FAQ가 많아질수록 챗봇이 FAQ를 고르는 시간이 조금 늘고 AI 비용이 늘어납니다. 수백 건에서도 문제없이 동작하는 것을 확인했습니다. 승인 FAQ가 약 500건을 넘으면 응답 속도와 비용을 점검해야 하니 개발사에 문의하세요.", "info", size=16)

# ── 6장
section_slide(6, "운영 루틴과 운영 상태", "언제 무엇을 할까요?")

s = new_slide("운영 루틴: 언제 무엇을 할까요?", "6. 운영 루틴")
table(s, 0.5, 1.65, 12.3, [1.6, 8, 1.1], [
    ["주기", "할 일", "걸리는 시간"],
    ["매일", "대시보드에서 '미처리 미답변 질문'을 확인하고 FAQ로 만들거나 '처리 완료' 표시. **불만 접수** 대화가 있는지 확인", "5분"],
    ["파일이 바뀔 때", "개정판 올리기 → '재검토 필요' 확인 → 각 FAQ를 새 파일과 비교해 고치거나 '변경 없음' 처리", "그때그때"],
    ["매주", "많이 묻는 질문 Top 10을 보고 담당자 이관이 많은 질문을 FAQ로 만들기. 답변 성공률이 오르는지 보기. '답변 완료' 몇 건을 골라 정확한지 확인. 고객이 쓴 말을 '다른 표현'에 추가", "30분"],
    ["매달", "오래되거나 중복된 FAQ 정리(초안으로 내리기). 출처 링크가 살아 있는지 확인. 대화 기록에 개인정보가 없는지 점검. 관리자 목록에서 퇴사자 빼기", "1시간"],
    ["제도·비용·기간이 바뀔 때", "해당 FAQ를 바로 고치고 저장(즉시 반영). 옛 내용으로 답한 대화가 없는지 대화 로그 확인", "그때그때"],
], size=16, max_h=5.1)

s = new_slide("운영 상태 화면: 챗봇이 건강한지 확인", "6. 운영 루틴")
textbox(s, 0.6, 1.6, 12.1, 0.9, ["위쪽 메뉴 **운영 상태**를 열 때마다 시스템이 **데이터 저장소와 AI에 실제로 연결**해 보고, 최근 1시간 오류·오늘 받은 질문 수·승인된 FAQ 수·처리 대기 중인 미답변 질문 수를 보여 줍니다."], size=18)
card(s, 0.6, 2.7, 3.9, 2.3, "정상 (초록)", ["문제가 없습니다"], fill=RGBColor(0xD1, 0xFA, 0xE5), edge=RGBColor(0x05, 0x96, 0x69), size=17)
card(s, 4.7, 2.7, 3.9, 2.3, "주의 (노랑)", ["승인된 FAQ가 0건", "오늘 질문 수가 하루 상한의 80%", "미답변 질문 30건 이상"], fill=AMBER_L, edge=AMBER, size=15)
card(s, 8.8, 2.7, 3.9, 2.3, "긴급 (빨강)", ["AI 연결 실패 (AI 사용 크레딧이 떨어졌거나 키 문제)", "데이터 저장소 연결 실패", "오류가 갑자기 많음 · 하루 상한 도달"], fill=RED_L, edge=RED, size=14)
callout(s, 0.6, 5.25, 12.1, 1.55, ["**'긴급'이 보이면 챗봇이 모든 질문을 \"담당자 확인\"으로 처리하고 있을 수 있습니다.** 바로 개발사에 알려 주세요.", "하루 한 번(오전 9시) 자동 점검이 돌고, 문제가 있으면 관리자 이메일로 알림 메일이 갑니다. 스팸함에 들어가면 '스팸 아님'으로 표시해 두세요."], "warn", size=15)

# ── 7장
section_slide(7, "실습", "직접 해 보기")

s = new_slide("실습 시나리오", "7. 실습")
table(s, 0.5, 1.6, 12.3, [0.5, 6.2, 5.2], [
    ["#", "해 볼 것", "확인할 것"],
    ["1", "[실습] FAQ를 초안으로 만들기 → 챗봇에 질문 → 승인 → 다시 질문", "초안일 때는 답하지 않고, 승인하면 바로 답함"],
    ["2", "FAQ에 없는 질문하기 → 미답변 질문에 쌓였는지 보기 → 직접 등록으로 FAQ 만들기", "저장하면 그 미답변이 자동으로 '처리 완료'"],
    ["3", "실습 FAQ를 초안으로 내리기 → 챗봇 확인 → 체크 칸으로 여러 건 삭제", "승인된 FAQ는 삭제되지 않음"],
    ["4", "파일 1개로 'FAQ 초안 만들기' → 검수하기 → 원문과 비교 → 1건 승인·1건 보류", "승인한 사람의 이메일이 남음"],
    ["5", "'보완 필요'인 비용 FAQ를 열어 답변 쓰기 → 승인", "'담당자 입력 필요'를 해제해야 승인됨"],
    ["6", "대화 로그에서 방금 나눈 대화를 찾아 근거 FAQ로 이동", "답이 이상하면 근거 FAQ를 고치면 됨"],
    ["7", "운영 상태 화면 열어 보기", "정상·주의·긴급의 의미 이해"],
], size=14, max_h=4.9, first_bold=False)
callout(s, 0.5, 6.35, 12.3, 0.55, "실습이 끝나면 [실습]이 붙은 FAQ를 모두 초안으로 되돌리거나 삭제합니다.", "warn", size=15)

# ── 8장
section_slide(8, "문제가 생겼을 때", "이럴 땐 이렇게 하세요")

s = new_slide("이럴 땐 이렇게 하세요", "8. 문제가 생겼을 때")
table(s, 0.5, 1.6, 12.3, [3.6, 6.6], [
    ["이런 일이 생기면", "이렇게 확인하세요"],
    ["FAQ를 저장했는데 챗봇이 예전 답을 함", "상태가 **승인**인지, '담당자 입력 필요'가 꺼져 있는지 확인. 고객 말투가 FAQ와 많이 다르면 **다른 표현**에 추가"],
    ["승인이 안 됨", "답변이 비어 있거나 '담당자 입력 필요'가 켜져 있음 (이유가 화면에 표시됨)"],
    ["챗봇이 엉뚱한 FAQ로 답함", "대화 로그에서 근거로 쓴 FAQ를 확인 → 질문을 더 구체적으로(인증 이름 포함) 고치거나 초안으로 내리기"],
    ["비용·기간이 안 나옴", "해당 FAQ가 '보완 필요'(답변이 비어 있음) 상태. 확정된 값을 쓰고 승인"],
    ["파일을 올렸는데 초안이 안 만들어짐", "상태가 **추출 완료**인지, 분류가 '기타'가 아닌지 확인. 스캔한 PDF는 글자가 있는 원본으로 다시 올리기"],
    ["챗봇이 계속 \"일시적인 문제로 답변을 드리지 못했습니다\"라고 함", "**운영 상태** 화면 확인 → '긴급'이면 개발사에 바로 연락"],
    ["고객이 \"잠시 후 다시 시도해 주세요\"를 봄", "짧은 시간에 요청이 너무 몰려 잠시 제한된 것. 잠시 뒤 정상으로 돌아옴. 계속되면 개발사에 문의"],
], size=14, max_h=5.2)

s = new_slide("알아 두면 좋은 현재 한계와 연락처", "8. 문제가 생겼을 때")
card(s, 0.6, 1.7, 6.1, 4.2, "알아 두면 좋은 현재 한계", [
    "비슷한 제품의 인증 사례(인증 이력)는 관리자 화면에서 입력할 수 없고, 개발사가 직접 넣습니다",
    "관리자를 추가하거나 빼려면 개발사에 요청해야 합니다",
    "챗봇은 한국어 질문만 지원합니다 (외국어는 담당자에게 연결)",
], size=16)
card(s, 6.9, 1.7, 5.8, 4.2, "연락처", ["전화 02-858-4321", "이메일 igsc@igsc.kr", "상담 신청 igsc.kr/contact/qna", "챗봇 igsc-chatbot.vercel.app/chat", "관리자 igsc-chatbot.vercel.app/admin"], size=16)
callout(s, 0.6, 6.1, 12.1, 0.75, "자세한 내용은 「IGSC 챗봇 사용자 교육 매뉴얼」 PDF를 참고하세요.", "info", size=17)

# 마지막
page[0] += 1
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, 13.333, 7.5, TEAL, shape=MSO_SHAPE.RECTANGLE)
textbox(s, 0.8, 2.8, 11.8, 1.2, [("감사합니다", {"bold": True})], size=48, color=WHITE)
textbox(s, 0.8, 4.1, 11.8, 1.0, ["질문과 의견을 자유롭게 말씀해 주세요."], size=22, color=RGBColor(0xCC, 0xFB, 0xF1))

prs.save(str(OUT))
print("생성:", OUT.name.encode("utf-8", "replace").decode("utf-8", "replace"), f"({page[0]}장)")
for w in warnings:
    print("경고:", w)
